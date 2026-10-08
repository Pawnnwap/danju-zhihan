#!/usr/bin/env python3
"""生成参赛 PPT：丹橘知旱——丹江口库区柑橘园气候决策助手"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

GREEN = RGBColor(0x0E, 0x52, 0x33)
GREEN2 = RGBColor(0x1A, 0x7A, 0x4A)
INK = RGBColor(0x1F, 0x2D, 0x27)
MUTED = RGBColor(0x6B, 0x7A, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background()
    r.shadow.inherit = False
    return s


def tb(slide, x, y, w, h, text, size=18, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        for run in p.runs:
            run.font.size = Pt(size); run.font.color.rgb = color; run.font.bold = bold
        if align:
            p.alignment = align
    return box


def header(slide, num, title):
    tb(slide, 0.6, 0.35, 1.2, 0.8, num, size=40, color=GREEN2, bold=True)
    tb(slide, 1.5, 0.52, 11, 0.7, title, size=28, color=GREEN, bold=True)
    bar = slide.shapes.add_shape(1, Inches(0.65), Inches(1.28), Inches(2.2), Inches(0.06))
    bar.fill.solid(); bar.fill.fore_color.rgb = GREEN2; bar.line.fill.background(); bar.shadow.inherit = False


# S1 封面
s = add_slide(GREEN)
tb(s, 1, 2.1, 11.3, 1.2, "丹橘知旱", size=60, color=WHITE, bold=True)
tb(s, 1, 3.4, 11.3, 0.7, "丹江口库区柑橘园气候决策助手", size=28, color=WHITE)
tb(s, 1, 4.25, 11.3, 0.5, "One Place · One Climate Problem · One AI Solution", size=16, color=RGBColor(0xCF, 0xE6, 0xD8))
tb(s, 1, 6.3, 11.3, 0.9, "AI4Climate 全球青年气候解决方案挑战赛 · 团队 沧浪舟\n开放选题 / 气候适应与韧性", size=14, color=RGBColor(0xCF, 0xE6, 0xD8))

# S2 真实问题
s = add_slide()
header(s, "01", "真实气候问题：极端旱涝下的库区柑橘")
tb(s, 0.7, 1.7, 12, 1.2,
   "丹江口水库——南水北调中线核心水源区；库周 30 万亩柑橘、约 10 万果农。\n气候变化已让「靠天吃饭」的损失常态化：", size=19)
for i, (y, t) in enumerate([
    ("2.95", "●  2022 年长江流域 1961 年来最强高温干旱：十堰竹山 44.6℃ 刷新湖北纪录；气象局首次派飞机在丹江口流域人工增雨"),
    ("3.75", "●  参考园块（凉水河镇）当年夏季 ≥35℃ 达 30 天、极端最高 40.9℃（ERA5-Land 复算）"),
    ("4.55", "●  产区研究：2008—2023 年 38 次气象灾害，干旱致花期落花率增加 32%—45%；越冬冻害频发"),
    ("5.35", "●  2025 年汉江 64 年来最密集秋汛、7 次编号洪水——旱涝急转成为新常态"),
]):
    tb(s, 0.7, float(y), 12, 0.7, t, size=16)
tb(s, 0.7, 6.35, 12, 0.7, "痛点：果农凭经验应对，缺少「我这块园 × 当前物候 × 未来几天」的决策提示，灾后补救远贵于灾前防范。",
   size=16, color=GREEN, bold=True)

# S3 方案
s = add_slide()
header(s, "02", "方案：从公开气象数据到可执行农事建议")
steps = [
    ("公开气象数据", "Open-Meteo\nERA5-Land 再分析\n+ ECMWF 预报\n（免密钥·浏览器直连）"),
    ("物候期识别", "萌芽/开花/幼果\n膨大/成熟/越冬\n六期规则表"),
    ("四类风险引擎", "高温热害 · 干旱\n花期低温 · 越冬冻害\n分级 + 触发依据"),
    ("农事建议", "风险 × 物候期\n动作清单\n标注规则来源"),
]
for i, (t, d) in enumerate(steps):
    x = 0.7 + i * 3.15
    card = s.shapes.add_shape(1, Inches(x), Inches(2.0), Inches(2.85), Inches(2.6))
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0xEF, 0xF6, 0xF1); card.line.color.rgb = GREEN2
    card.shadow.inherit = False
    tf = card.text_frame; tf.word_wrap = True
    tf.text = t
    for run in tf.paragraphs[0].runs:
        run.font.size = Pt(19); run.font.bold = True; run.font.color.rgb = GREEN
    p = tf.add_paragraph(); p.text = d
    for run in p.runs:
        run.font.size = Pt(13); run.font.color.rgb = INK
    if i < 3:
        ar = s.shapes.add_shape(13, Inches(x + 2.88), Inches(3.05), Inches(0.25), Inches(0.45))
        ar.fill.solid(); ar.fill.fore_color.rgb = GREEN2; ar.line.fill.background(); ar.shadow.inherit = False
tb(s, 0.7, 5.1, 12, 0.6, "在线体验：https://pawnnwap.github.io/danju-zhihan/      代码开源：https://github.com/Pawnnwap/danju-zhihan", size=15, color=GREEN2, bold=True)
tb(s, 0.7, 5.8, 12, 1.2,
   "手机浏览器打开即用：点选园块（或输入自定坐标）→ 当前物候与四类风险看板 → 90 天风险日历 → 农事建议 → 历史回放 → 问答。", size=15)

# S4 AI 任务
s = add_slide()
header(s, "03", "AI 承担的具体任务：规则算得准，大模型讲得懂")
tb(s, 0.7, 1.75, 5.8, 0.5, "① 确定性风险引擎（核心）", size=20, color=GREEN, bold=True)
tb(s, 0.7, 2.3, 5.8, 3.4,
   "●  阈值 + 连续日数 + 物候期组合判断，逐日输出\n     等级与触发依据\n●  全部阈值显式可审计、可复算、可回测\n●  同一套规则有 Python 独立实现交叉验证", size=16)
tb(s, 6.9, 1.75, 5.8, 0.5, "② DeepSeek 大模型解释层（可选接入）", size=20, color=GREEN, bold=True)
tb(s, 6.9, 2.3, 5.8, 3.4,
   "●  以风险数据 + 农技规则库为上下文，生成个性化\n     建议与自然语言追问\n●  回答强制引用数据日期，抑制幻觉\n●  LLM 只做转译与解释，不做数值预测", size=16)
tb(s, 0.7, 5.9, 12, 0.9,
   "为什么不是「为了 AI 而 AI」：没有大模型时产品依然完整可用；大模型解决的是「同一份数据对不同用户的不同表述 + 追问」这一具体负担。",
   size=15, color=MUTED)

# S5 验证
s = add_slide()
header(s, "04", "验证：历史回测已通过，90 天试点已设计")
card = s.shapes.add_shape(1, Inches(0.7), Inches(1.8), Inches(12), Inches(2.3))
card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0xEF, 0xF6, 0xF1); card.line.color.rgb = GREEN2
card.shadow.inherit = False
tf = card.text_frame; tf.word_wrap = True
tf.text = "历史回测（2022 年热害 · 凉水河镇 · 真实数据）"
for run in tf.paragraphs[0].runs:
    run.font.size = Pt(20); run.font.bold = True; run.font.color.rgb = GREEN
for line in ["≥35℃ 高温日 30 天    热害『重』级 13 天    极端最高 40.9℃（08-07）",
             "首个『重』级提示 2022-08-04 —— 比峰值提前 3 天，覆盖灾前防护窗口",
             "页面内置 2022 热害 / 2025 秋汛 / 2021 冬寒潮三个回放场景，可自助复验"]:
    p = tf.add_paragraph(); p.text = "●  " + line
    for run in p.runs:
        run.font.size = Pt(16); run.font.color.rgb = INK
tb(s, 0.7, 4.5, 12, 2.2,
   "90 天试点计划：\n●  对象：库区 1—2 家合作社 / 乡镇农技员（30—100 亩）\n●  指标：建议采纳率 ≥50% ｜ 提示时效提前 ≥1 天 ｜ 满意度 ≥3.8/5 ｜ 续用意愿 ≥80%\n●  迭代：阈值与物候表每两周校准；高频追问沉淀为规则条目", size=17)

# S6 市场与安全
s = add_slide()
header(s, "05", "可持续 · 开放 · 安全边界")
tb(s, 0.7, 1.8, 5.8, 0.5, "市场可持续", size=20, color=GREEN, bold=True)
tb(s, 0.7, 2.35, 5.8, 2.6,
   "●  合作社年费（数百元级/园块）\n●  气象数据授权运营场景分成\n     （湖北已发首批授权）\n●  「气象+保险」联动：风险日志作为\n     投保/理赔辅助凭证\n●  小农户基础功能永久免费", size=16)
tb(s, 6.9, 1.8, 5.8, 0.5, "扩展与伦理", size=20, color=GREEN, bold=True)
tb(s, 6.9, 2.35, 5.8, 2.6,
   "●  换物候表 + 坐标即可复制到其他柑橘产区\n●  同一引擎可扩展至茶园等库区作物\n●  代码开源，数据全部公开来源\n●  不采集个人身份信息，可随时退出", size=16)
tb(s, 0.7, 5.3, 12, 1.3,
   "安全边界：风险等级用于辅助安排农事，不构成灾害预警，不替代官方气象预警与专业农技指导；\n再分析数据有不确定性，重要决策结合当地气象部门信息。", size=16, color=RGBColor(0xC0, 0x39, 0x2B))

# S7 团队
s = add_slide(GREEN)
tb(s, 1, 2.4, 11.3, 1, "沧浪舟", size=44, color=WHITE, bold=True)
tb(s, 1, 3.6, 11.3, 0.6, "个人参赛 · 2026 年 10 月", size=18, color=RGBColor(0xCF, 0xE6, 0xD8))
tb(s, 1, 4.6, 11.3, 1.2,
   "把一个真实问题带来，让 AI 真正去解决它。\n一个地方。一个气候问题。一个 AI 解决方案。", size=20, color=WHITE)
tb(s, 1, 6.3, 11.3, 0.5, "ai4climate.world ｜ pawnnwap.github.io/danju-zhihan ｜ deepseek.club/topic/5043", size=13, color=RGBColor(0xCF, 0xE6, 0xD8))

prs.save("丹橘知旱-参赛PPT.pptx")
print("saved 丹橘知旱-参赛PPT.pptx")
